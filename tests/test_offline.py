"""Offline test suite — no network, no real credentials. Run: pytest -q"""

import os

os.environ.setdefault("SUPABASE_URL", "https://dummy.supabase.co")
os.environ.setdefault("SUPABASE_SERVICE_KEY", "dummy")
os.environ.setdefault("TELEGRAM_BOT_TOKEN", "dummy")
os.environ.setdefault("TELEGRAM_OWNER_CHAT_ID", "1")
os.environ.setdefault("OPENROUTER_API_KEY", "dummy")
os.environ.setdefault("GUMROAD_ACCESS_TOKEN", "dummy")


def test_scoring_formula():
    from app.scanner.scoring import competition_penalty

    assert competition_penalty(500) == 1.0
    assert competition_penalty(3000) == 0.6
    assert competition_penalty(10000) == 0.3
    assert competition_penalty(50000) == 0.1
    assert competition_penalty(None) == 0.5


def test_candidates_relevance_and_dedupe():
    from app.config import NICHES
    from app.scanner.scoring import build_candidates, finalize_scores

    terms = NICHES["spreadsheets"]["relevance_terms"]
    cands = build_candidates(
        {"budget spreadsheet": 80.0, "dog collar": 90.0},
        {"budget spreadsheet": ["wedding budget spreadsheet", "dog leash"]},
        terms,
    )
    kws = {c.keyword for c in cands}
    assert "dog collar" not in kws and "wedding budget spreadsheet" in kws
    finalize_scores(cands)


def test_llm_json_extraction():
    from app.llm import extract_json

    assert extract_json('```json\n{"a": 1}\n```') == {"a": 1}
    assert extract_json('prose {"a": 1} prose') == {"a": 1}


def test_listing_copy_constraints():
    from app.generator import _finalize_copy

    c = _finalize_copy({"title": "T" * 300, "tags": ["a", "a"], "price_usd": 3}, "wedding budget")
    assert len(c["title"]) == 140
    assert len(c["tags"]) == 13 == len(set(c["tags"]))
    assert c["price_usd"] == 9.0
    assert _finalize_copy({"price_usd": 99}, "x")["price_usd"] == 19.0


def test_xlsx_builder(tmp_path):
    from openpyxl import load_workbook

    from app.xlsx_builder import build_xlsx

    spec = {
        "product_name": "Test",
        "theme": {"primary": "112233"},
        "sheets": [{
            "name": "Data",
            "columns": [{"header": "A", "format": "currency"}, {"header": "B"}],
            "rows": [["x", "=A4*2"]],
            "totals_row": ["T", "=SUM(B4:B4)"],
        }],
        "how_to_use": ["Step"],
    }
    out = str(tmp_path / "t.xlsx")
    build_xlsx(spec, out)
    wb = load_workbook(out)
    ws = wb["Data"]
    assert ws.freeze_panes == "A4"
    assert ws["B4"].value == "=A4*2"
    assert ws["B5"].value == "=SUM(B4:B4)"


def test_cover_render(tmp_path):
    from PIL import Image

    from app.cover import render_cover

    spec = {"product_name": "Wedding Budget Pro", "tagline": "Plan it all",
            "theme": {"primary": "6B4E71", "secondary": "F3EDF5", "accent": "D4A947"},
            "sheets": [{"columns": [{"header": "Item"}, {"header": "Cost"}],
                        "rows": [["Venue", 5000], ["Food", "=B4*2"]]}]}
    out = str(tmp_path / "cover.png")
    render_cover(spec, out)
    img = Image.open(out)
    assert img.size == (1280, 720)


def test_calibration_sample_guards():
    from app.brain import calibrate

    def product(source, sales):
        return {"total_sales": sales, "total_revenue": 10.0 * sales, "suggested_price_usd": 12,
                "opportunities": {"source": source, "competition_count": 100, "niche": "n"}}

    rows = [product("etsy_autocomplete", 1) for _ in range(6)] + [product("google_trends", 5)]
    lines, evidence = calibrate(rows)
    text = " ".join(lines)
    assert "etsy_autocomplete" in text
    assert "google_trends" not in text  # n=1 bucket must stay silent
    assert evidence["buckets"]["source:google_trends"]["n"] == 1


def test_competitor_ldjson_parser():
    from app.scanner.competitors import _listings_from_ldjson

    html = '''<script type="application/ld+json">
    {"@type":"ItemList","itemListElement":[
      {"item":{"name":"Tracker Pro","offers":{"price":"12.99"},
       "aggregateRating":{"ratingValue":"4.9","reviewCount":"88"}}}]}</script>'''
    listings = _listings_from_ldjson(html)
    assert listings == [{"title": "Tracker Pro", "price": "12.99", "rating": "4.9", "reviews": "88"}]


def test_seasonal_seeds():
    from app.config import seasonal_seeds

    july = seasonal_seeds("spreadsheets", 7)
    assert "back to school budget" in july          # current month
    assert "back to school checklist" in july       # next month
    assert seasonal_seeds("unknown_niche", 7) == []


def test_niche_activation():
    import pytest

    from app.config import active_niches

    os.environ["TRENDMILL_NICHES"] = "trader_tools"
    assert list(active_niches()) == ["trader_tools"]
    os.environ["TRENDMILL_NICHES"] = "nope"
    with pytest.raises(RuntimeError):
        active_niches()
    os.environ.pop("TRENDMILL_NICHES")


def test_uploader_dispatch_error():
    import pytest

    from app.uploader import upload_product

    os.environ["UPLOAD_STRATEGY"] = "bogus"
    with pytest.raises(ValueError):
        upload_product({})
    os.environ.pop("UPLOAD_STRATEGY")


def test_learner_aggregation():
    from app.learner import aggregate_sales

    agg = aggregate_sales([
        {"product_id": "p", "price": 1000},
        {"product_id": "p", "price": 1000, "refunded": True},
    ])
    assert agg["p"] == {"sales_count": 1, "revenue_usd": 10.0}


def test_server_routes():
    from app.server import app

    paths = {r.path for r in app.routes}
    assert {"/health", "/telegram/webhook"} <= paths


def test_quality_structural_checks(tmp_path):
    from app.quality import structural_issues

    good_spec = {
        "sheets": [
            {"rows": [["a", "=B4"]] * 9, "totals_row": ["T", "=SUM(B4:B12)"]},
            {"rows": [["x", 1]] * 8},
        ],
        "how_to_use": ["1", "2", "3", "4", "5"],
    }
    good_copy = {"description": "d" * 400}
    f = tmp_path / "p.xlsx"
    f.write_bytes(b"0" * 5000)
    assert structural_issues(good_spec, good_copy, str(f)) == []

    bad = structural_issues({"sheets": [{"rows": [["a"]]}], "how_to_use": ["1"]},
                            {"description": "short"}, str(tmp_path / "missing.xlsx"))
    assert len(bad) >= 4  # sheets, rows, formulas, steps, description, file


def test_extra_niches_env():
    import json

    import pytest

    from app.config import active_niches

    os.environ["TRENDMILL_EXTRA_NICHES"] = json.dumps({
        "hebrew_il": {"seeds": ["s"], "relevance_terms": ["t"]}
    })
    niches = active_niches()
    assert "hebrew_il" in niches and niches["hebrew_il"]["label"] == "hebrew_il"

    os.environ["TRENDMILL_EXTRA_NICHES"] = json.dumps({"bad": {"seeds": []}})
    with pytest.raises(RuntimeError):
        active_niches()
    os.environ.pop("TRENDMILL_EXTRA_NICHES")


def test_foresight_seasonal_profile():
    import pandas as pd

    from app.scanner.foresight import seasonal_profile

    # 3+ calendar years, ramp weeks 12-14, peak weeks 15-20, quiet otherwise
    idx = pd.date_range("2022-01-03", periods=170, freq="W-MON")
    values = []
    for ts in idx:
        w = ts.isocalendar().week
        values.append(80 if 15 <= w <= 20 else 40 if 12 <= w <= 14 else 20)
    profile = seasonal_profile(pd.Series(values, index=idx))
    assert profile is not None
    assert profile["peak_week"] == 15 or 15 <= profile["peak_week"] <= 20
    assert profile["rise_week"] == 12          # caught the ramp start
    assert profile["strength"] >= 1.5

    # Guard: flat series has no seasonality
    flat = pd.Series([30] * 170, index=idx)
    assert seasonal_profile(flat) is None

    # Guard: too little history
    short = pd.Series(values[:40], index=idx[:40])
    assert seasonal_profile(short) is None


def test_foresight_lead_lag():
    import numpy as np
    import pandas as pd

    from app.scanner.foresight import lead_lag

    rng = np.random.default_rng(42)
    idx = pd.date_range("2021-01-04", periods=200, freq="W-MON")
    driver = 50 + 20 * np.sin(np.arange(200) / 5) + rng.normal(0, 1.5, 200)
    lag = 3
    follower = np.roll(driver, lag) + rng.normal(0, 1.0, 200)

    link = lead_lag(pd.Series(driver, index=idx), pd.Series(follower, index=idx))
    assert link is not None
    assert link["lag_weeks"] == lag            # recovered the true 3-week lag
    assert link["correlation"] >= 0.45

    # Guard: pure noise pair must not produce a link
    noise_a = pd.Series(50 + rng.normal(0, 5, 200), index=idx)
    noise_b = pd.Series(50 + rng.normal(0, 5, 200), index=idx)
    assert lead_lag(noise_a, noise_b) is None
